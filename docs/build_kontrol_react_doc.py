#!/usr/bin/env python3
"""
Контрольные вопросы (React, props/state, компонент/элемент, useState/useEffect, жизненный цикл, Vite/Babel):
скриншоты слайдов лекции 8 + * где есть; без * — только там, где для ответа не хватало * (умеренно).
Лекция 9 здесь не используется — её темы не в перечне контрольных вопросов.
Запуск: python3 docs/build_kontrol_react_doc.py
"""
from __future__ import annotations

from pathlib import Path

import fitz
from docx import Document
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
from docx.shared import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
OUT_DOC = ROOT / "docs" / "Kontrolnye_React_Lekcii_8_9.docx"
SLIDES_DIR = ROOT / "docs" / "_kontrol_react_slides"

L8 = ROOT / "Lecture_8_React_Introduction.pdf"

ZOOM = 2.15
IMG_WIDTH = Inches(6.35)

# Порядок = порядок контрольных тем. Пометка: * = в заголовке слайда на презентации есть звёздочка.
SLIDE_EXPORTS = [
    (L8, 5, "l8_p05.png", "React (слайд без *, добавлен к ответу «что такое React»)"),
    (L8, 6, "l8_p06.png", "JSX (без *, к паре «компонент — элемент»)"),
    (L8, 7, "l8_p07.png", "Компоненты *"),
    (L8, 8, "l8_p08.png", "Props (без *, отдельного * про только props нет)"),
    (L8, 9, "l8_p09.png", "Состояние *"),
    (L8, 10, "l8_p10.png", "Хуки. useState *"),
    (L8, 12, "l8_p12.png", "Жизненный цикл приложения *"),
    (L8, 13, "l8_p13.png", "Методы жизненного цикла *"),
    (L8, 14, "l8_p14.png", "useEffect *"),
    (L8, 18, "l8_p18.png", "Babel (без *, для * нет отдельного слайда)"),
    (L8, 23, "l8_p23.png", "Vite + React. Основные файлы *"),
    (L8, 24, "l8_p24.png", "Vite + React. App.tsx *"),
    (L8, 26, "l8_p26.png", "Роутинг *"),
]


def export_slide(pdf_path: Path, page_1based: int, out_name: str) -> Path:
    SLIDES_DIR.mkdir(parents=True, exist_ok=True)
    doc = fitz.open(pdf_path)
    page = doc[page_1based - 1]
    mat = fitz.Matrix(ZOOM, ZOOM)
    pix = page.get_pixmap(matrix=mat, alpha=False)
    out = SLIDES_DIR / out_name
    pix.save(str(out))
    doc.close()
    return out


def add_code(doc: Document, path_hint: str, text: str) -> None:
    p = doc.add_paragraph()
    r = p.add_run(f"{path_hint}\n{text.rstrip()}")
    r.font.name = "Consolas"
    r.font.size = Pt(8)
    p.paragraph_format.left_indent = Inches(0.15)
    p.paragraph_format.space_after = Pt(4)


def add_slide_block(doc: Document, caption: str, png: Path, note: str) -> None:
    doc.add_paragraph(caption, style="Intense Quote")
    doc.add_picture(str(png), width=IMG_WIDTH)
    doc.add_paragraph(note)
    doc.add_paragraph()


def main() -> None:
    by_name: dict[str, Path] = {}
    for pdf, num, name, _hint in SLIDE_EXPORTS:
        by_name[name] = export_slide(pdf, num, name)

    doc = Document()
    h = doc.add_heading("Контрольные вопросы: React (лекция 8)", 0)
    h.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    doc.add_paragraph(
        "Темы: React; props и состояние; компонент и элемент; useState и useEffect; жизненный цикл; Vite и Babel. "
        "Скриншоты в основном со слайдов с *; где отдельного * не было (React, JSX, Props, Babel), добавлен один обычный слайд на тему — без лишних страниц. "
        "Фрагменты кода — актуальный new_front: каталог типов аккумуляторов (battery_life_types), маршруты BATTERY_TYPES / BATTERY_TYPE, интерфейс BatteryTypeMock; сценарий гостя без заявки."
    )

    # --- 1. React ---
    doc.add_heading("1. React", level=1)
    add_slide_block(
        doc,
        "Лекция 8 — слайд «React» (без * в заголовке)",
        by_name["l8_p05.png"],
        "React — библиотека для работы с виртуальным DOM. Документация: reactjs.org/docs/getting-started.html.",
    )

    # --- 2. Компонент и элемент ---
    doc.add_heading("2. Компонент и React-элемент", level=1)
    add_slide_block(
        doc,
        "Лекция 8 — «JSX» (без *)",
        by_name["l8_p06.png"],
        "В компонентах используется JSX/TSX. После компиляции JSX получаются простые объекты — React-элементы. "
        "Пропсы в DOM-стиле — camelCase (tabIndex, className).",
    )
    add_slide_block(
        doc,
        "Лекция 8 — «Компоненты *»",
        by_name["l8_p07.png"],
        "Компонент — переиспользуемая единица, возвращающая элементы для страницы. Три акцента лабы: отрисовка, состояние, события. "
        "Функциональные и классовые компоненты.",
    )
    add_code(
        doc,
        "new_front/src/main.tsx — элемент <App /> в дереве",
        """ReactDOM.createRoot(document.getElementById("root")!).render(
  <React.StrictMode>
    <App />
  </React.StrictMode>,
);""",
    )
    add_code(
        doc,
        "new_front/src/routePaths.ts и App.tsx — компонент App и маршруты каталога",
        """// routePaths.ts
export const ROUTES = {
  BATTERY_TYPES: "/",
  BATTERY_TYPE: "/battery/:id",
} as const;

// App.tsx (фрагмент)
import BatteryTypesPage from "./pages/BatteryTypesPage/BatteryTypesPage";
import BatteryTypePage from "./pages/BatteryTypePage/BatteryTypePage";

<Route path={ROUTES.BATTERY_TYPES} element={<BatteryTypesPage />} />
<Route path={ROUTES.BATTERY_TYPE} element={<BatteryTypePage />} />""",
    )

    # --- 3. Props и состояние ---
    doc.add_heading("3. Props и состояние", level=1)
    add_slide_block(
        doc,
        "Лекция 8 — «Props» (без *)",
        by_name["l8_p08.png"],
        "Props — входные данные от родителя к потомку. props.children — содержимое между тегами. В классах: this.props.children.",
    )
    add_slide_block(
        doc,
        "Лекция 8 — «Состояние *»",
        by_name["l8_p09.png"],
        "State — когда данные меняются со временем. Состояние управляет компонентом, props передают информацию. "
        "Обычная переменная не даёт перерисовку; смена state почти всегда вызывает render.",
    )
    add_code(
        doc,
        "new_front: state в родителе, props в BatteryTypesFilterBar",
        """// BatteryTypesPage.tsx
const [filters, setFilters] = useState<BatteryFilters>(initialFilters);
<BatteryTypesFilterBar
  query={filters.title}
  onQueryChange={(q) => setFilters((f) => ({ ...f, title: q }))}
  onSearch={applyFilters}
/>

// BatteryTypesFilterBar.tsx — только props
export default function BatteryTypesFilterBar({ query, onQueryChange, onSearch }: BatteryTypesFilterBarProps) { /* ... */ }""",
    )

    # --- 4. useState и useEffect ---
    doc.add_heading("4. useState и useEffect", level=1)
    add_slide_block(
        doc,
        "Лекция 8 — «Хуки. useState *»",
        by_name["l8_p10.png"],
        "Пример: onClick меняет счётчик; значение state в JSX; при изменении state — новая отрисовка компонента.",
    )
    add_slide_block(
        doc,
        "Лекция 8 — «useEffect *»",
        by_name["l8_p14.png"],
        "Идеи componentDidMount / componentDidUpdate / componentWillUnmount. useEffect — не весь жизненный цикл, а три «зелёных» блока со схемы; цикл шире (props, render…).",
    )
    add_code(
        doc,
        "new_front/src/pages/BatteryTypePage/BatteryTypePage.tsx — useState + useEffect по id из URL",
        """const [battery, setBattery] = useState<BatteryTypeMock | null>(null);
const { id } = useParams();

useEffect(() => {
  if (!id) {
    setBattery(null);
    return;
  }
  setMediaError(false);
  const n = Number(id);
  const resolved =
    getMockBattery(n) ?? BATTERIES_MOCK.find((b) => b.battery_id === n) ?? null;
  setBattery(resolved);
}, [id]);""",
    )
    add_code(
        doc,
        "new_front/src/components/BatteryTypeCard/BatteryTypeCard.tsx — useEffect при смене props",
        """useEffect(() => {
  setImageError(false);
  setImageUrl(photoSrc(battery.photo_url, false));
}, [battery.photo_url]);""",
    )

    # --- 5. Жизненный цикл ---
    doc.add_heading("5. Жизненный цикл компонента", level=1)
    add_slide_block(
        doc,
        "Лекция 8 — «Жизненный цикл приложения *»",
        by_name["l8_p12.png"],
        "(1) Монтирование: getDerivedStateFromProps, render, JSX, попадание в DOM. (2) Обновление: при смене state или props. "
        "(3) Размонтирование: componentWillUnmount перед удалением из DOM.",
    )
    add_slide_block(
        doc,
        "Лекция 8 — «Методы жизненного цикла *»",
        by_name["l8_p13.png"],
        "Полная схема методов для классового компонента — на скриншоте.",
    )

    # --- 6. Vite и Babel ---
    doc.add_heading("6. Vite и Babel", level=1)
    add_slide_block(
        doc,
        "Лекция 8 — «Babel» (без *)",
        by_name["l8_p18.png"],
        "Babel — транспайлер; например ES6 (ES2015) → более старый ES5.",
    )
    add_slide_block(
        doc,
        "Лекция 8 — «Vite + React. Основные файлы *»",
        by_name["l8_p23.png"],
        "index.html + main: точка входа дерева компонентов, всё в App. Vite — альтернатива CRA, сборщик модулей JS.",
    )
    add_slide_block(
        doc,
        "Лекция 8 — «Vite + React. App.tsx *»",
        by_name["l8_p24.png"],
        "App — главный компонент; шаблон с кнопкой и состоянием; дальше — разбиение на компоненты.",
    )
    add_slide_block(
        doc,
        "Лекция 8 — «Роутинг *» (к шаблону Vite/React в курсе)",
        by_name["l8_p26.png"],
        "Роутер в main, переходы по URL; можно упростить App до элементов по маршрутам.",
    )
    add_code(
        doc,
        "new_front — Vite (index.html, vite.config.ts)",
        """// index.html
<div id="root"></div>
<script type="module" src="/src/main.tsx"></script>

// vite.config.ts
export default defineConfig({
  plugins: [react()],
  server: {
    port: 3000,
    proxy: { "/api": { target: "http://localhost:8080", changeOrigin: true } },
  },
});""",
    )
    add_code(
        doc,
        "new_front/src/modules/batteryApi.ts — тип mock-записи каталога (не «Service»)",
        """export interface BatteryTypeMock {
  battery_id: number;
  title: string;
  short_description: string;
  // ... photo_url, video, capacity_mah, voltage_v и др.
}""",
    )

    doc.save(OUT_DOC)
    print(f"Written: {OUT_DOC}")


if __name__ == "__main__":
    main()
